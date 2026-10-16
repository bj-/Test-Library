#!/usr/bin/env python3
"""Manual accuracy check of LLM steps (stage 7, ADR 0005) as an Excel workbook.

Usage (from Api/):
  accuracy.py build --batch pilot    ideas/mined/<batch>.accuracy.xlsx: sample (evaluate.sample) + Claude's first
                                     pass (<batch>.accuracy.claude.yaml) + human marks kept from the existing workbook
  accuracy.py score --batch pilot    accuracy by section from the filled workbook
  accuracy.py build --batch a,b,c --name scale-01 --match-only
                                     one sample over several batches (sources); --match-only — matching only,
                                     --n — sample size; files are named by --name

The human fills «Верно (человек)» (да / нет / не знаю) and «Комментарий». Rows to check are marked
«обязательно» (Claude doubts or found a mismatch) and «выборочно» (random spot check of Claude's «совпадает»).
A row counts by the human mark if present (да / нет), otherwise — also for «не знаю» — by Claude's verdict
(ok -> да, wrong -> нет, doubt -> open).
"""
from __future__ import annotations

import argparse
import random
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from cluster import MINED, load_results, load_sources  # noqa: E402
from evaluate import SAMPLE_MATCH, SEED, sample  # noqa: E402
from library import load_yaml, setup_console  # noqa: E402

SPOT_PER_SECTION = 6
TARGET = 85.0
VERDICT = {"ok": "совпадает", "doubt": "сомнение", "wrong": "не совпадает"}
FILL = {"ok": "E2EFDA", "doubt": "FFF2CC", "wrong": "F8CBAD"}
HEAD = ["ID", "Раздел", "Баг", "Источник", "Решение LLM", "Паттерн / сопоставление", "Выдержка из бага",
        "Вывод Claude", "Почему", "Проверить", "Верно (человек)", "Комментарий"]
WIDTH = [6, 12, 26, 12, 22, 50, 60, 14, 50, 13, 14, 50]
MARKS = ("да", "нет", "не знаю")
HELP = [
    "Что проверять",
    "Вопрос не «нужна ли такая проверка» и не «баг ли это вообще», а «правильно ли LLM пересказала исходный баг».",
    "Открыть ссылку в колонке «Баг» и сравнить с колонками «Решение LLM» и «Паттерн / сопоставление»:",
    "• Триаж: верно ли решение «API» / «вне области» и класс корзины;",
    "• Пересказ (строки API): совпадает ли паттерн с багом — не перепутана суть, не додумано ожидание, которого нет;",
    "• Сопоставление: действительно ли баг относится к записи (covered — уже покрыт, extends — новый вариант записи).",
    "",
    "Как заполнять",
    "1. Фильтр по колонке «Проверить»: «обязательно» — Claude сомневается или нашёл расхождение; «выборочно» — "
    "случайная проверка строк, где Claude считает «совпадает».",
    "2. В «Верно (человек)» выбрать да / нет / не знаю; в «Комментарий» — почему (особенно для «нет»).",
    "3. Остальные строки можно не трогать: засчитывается вывод Claude (и при «не знаю» тоже). Отметка да/нет всегда важнее.",
    "4. Идеи новых проверок — тоже в «Комментарий»: Claude перенесёт их в Api/ideas/inbox.md.",
    "",
    f"Цель — не меньше {TARGET:.0f} % верных отдельно для триажа и для сопоставления (ADR 0005).",
    "Итог считает лист «Итог»; точный расчёт — python tools/mining/accuracy.py score --batch {name}.",
]


def batches(args) -> list[str]:
    return [b.strip() for b in args.batch.split(",") if b.strip()]


def name(args) -> str:
    return args.name or args.batch


def sample_many(names: list[str], n: int, match_only: bool) -> tuple[list[dict], list[dict]]:
    """Pilot (one batch, full sample) stays as evaluate.sample; otherwise one pool over all batches, seed SEED."""
    if len(names) == 1 and not match_only:
        return sample(names[0])
    rnd = random.Random(SEED)
    pool = [r for b in names for r in load_results(b) if r["verdict"] == "api" and r["match"]["type"] != "new"]
    return [], rnd.sample(pool, n)


def rows(args) -> list[dict]:
    names = batches(args)
    triage, matched = sample_many(names, args.n, args.match_only)
    claude = load_yaml(MINED / f"{name(args)}.accuracy.claude.yaml")
    sources = {k: v for b in names for k, v in load_sources(b).items()}
    out = []
    for prefix, section, recs in (("T", "Триаж", triage), ("M", "Сопоставление", matched)):
        for n, r in enumerate(recs, 1):
            key = f"{prefix}{n}"
            c = claude.get(key) or {}
            if r["verdict"] == "api":
                decision = "API"
                what = r["pattern"]["title"]
            else:
                decision = f"вне области: {r['oos_class']}"
                what = r["reason"]
            if section == "Сопоставление":
                m = r["match"]
                what = f"{m['type']} → {m['check_id']}\n{m.get('note') or ''}\nПаттерн: {r['pattern']['title']}"
            out.append({"key": key, "section": section, "bug": r["id"], "url": sources[r["id"]]["url"],
                        "source": sources[r["id"]]["source"], "decision": decision, "what": what,
                        "quote": c.get("quote", ""), "verdict": c.get("verdict"), "why": c.get("why", "")})
    return out


def check_marks(items: list[dict]) -> None:
    rnd = random.Random(SEED)
    for section in ("Триаж", "Сопоставление"):
        sec = [i for i in items if i["section"] == section]
        for i in sec:
            i["check"] = "обязательно" if i["verdict"] in ("doubt", "wrong") else ""
        ok = [i for i in sec if i["verdict"] == "ok"]
        for i in rnd.sample(ok, min(SPOT_PER_SECTION, len(ok))):
            i["check"] = "выборочно"


def old_marks(path: Path) -> dict[str, tuple]:
    if not path.exists():
        return {}
    ws = load_workbook(path)["Выборка"]
    return {r[0]: (r[10], r[11]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}


def cmd_build(args) -> int:
    path = MINED / f"{name(args)}.accuracy.xlsx"
    items = rows(args)
    missing = [i["key"] for i in items if not i["verdict"]]
    if missing:
        print("ОШИБКА: нет вывода Claude для", ", ".join(missing))
        return 1
    check_marks(items)
    marks = old_marks(path)
    if args.marks:  # one-off migration: KEY=mark|comment lines
        for line in Path(args.marks).read_text(encoding="utf-8").splitlines():
            if "=" in line:
                key, rest = line.split("=", 1)
                mark, _, comment = rest.partition("|")
                marks[key.strip()] = (mark.strip() or None, comment.strip() or None)

    wb = Workbook()
    ws = wb.active
    ws.title = "Выборка"
    ws.append(HEAD)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    for i in items:
        mark, comment = marks.get(i["key"], (None, None))
        ws.append([i["key"], i["section"], i["bug"], i["source"], i["decision"], i["what"], i["quote"],
                   VERDICT[i["verdict"]], i["why"], i["check"], mark, comment])
        row = ws.max_row
        ws.cell(row, 3).hyperlink = i["url"]
        ws.cell(row, 3).font = Font(color="0563C1", underline="single")
        ws.cell(row, 8).fill = PatternFill("solid", fgColor=FILL[i["verdict"]])
        if i["check"] == "обязательно":
            ws.cell(row, 10).font = Font(bold=True)
    for n, w in enumerate(WIDTH, 1):
        ws.column_dimensions[chr(64 + n)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col in "KL":
        for row in range(2, ws.max_row + 1):
            ws[f"{col}{row}"].fill = PatternFill("solid", fgColor="FFFFFF" if col == "L" else "FCE4D6")
    dv = DataValidation(type="list", formula1='"' + ",".join(MARKS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"K2:K{ws.max_row}")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:L{ws.max_row}"

    total = wb.create_sheet("Итог")
    total.append(["Раздел", "Строк", "Человек: да", "Человек: нет", "Человек: не знаю", "Claude: совпадает",
                  "Claude: сомнение", "Claude: не совпадает", "Верно, % (человек, иначе Claude; «не знаю» = Claude; сомнения не в расчёте)"])
    rng = "Выборка!$B:$B"

    def fb(r: int, verdict: str) -> str:  # rows without a human mark (empty or «не знаю») count by Claude
        return "+".join(f'COUNTIFS({rng},A{r},Выборка!$K:$K,"{m}",Выборка!$H:$H,"{verdict}")' for m in ("", "не знаю"))

    for r, section in enumerate(("Триаж", "Сопоставление"), 2):
        total.append([section, f'=COUNTIF({rng},A{r})',
                      f'=COUNTIFS({rng},A{r},Выборка!$K:$K,"да")', f'=COUNTIFS({rng},A{r},Выборка!$K:$K,"нет")',
                      f'=COUNTIFS({rng},A{r},Выборка!$K:$K,"не знаю")',
                      f'=COUNTIFS({rng},A{r},Выборка!$H:$H,"совпадает")',
                      f'=COUNTIFS({rng},A{r},Выборка!$H:$H,"сомнение")',
                      f'=COUNTIFS({rng},A{r},Выборка!$H:$H,"не совпадает")',
                      f'=IFERROR(100*(C{r}+{fb(r, "совпадает")})/(C{r}+D{r}+{fb(r, "совпадает")}+{fb(r, "не совпадает")}),"")'])
    total.append([])
    total.append([f"Цель — не меньше {TARGET:.0f} % в каждом разделе (ADR 0005)."])
    for cell in total[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHI", (16, 8, 12, 12, 14, 14, 14, 16, 30)):
        total.column_dimensions[col].width = w

    hp = wb.create_sheet("Как заполнять", 0)
    for line in HELP:
        hp.append([line.replace("{name}", name(args))])
    hp.column_dimensions["A"].width = 140
    for r in (1, 8):
        hp.cell(r, 1).font = Font(bold=True, size=12)
    wb.active = 1
    wb.save(path)
    c = Counter(i["check"] for i in items)
    print(f"{path}: строк {len(items)}, обязательно {c['обязательно']}, выборочно {c['выборочно']},"
          f" отметок человека {sum(1 for v in marks.values() if v[0])}")
    return 0


def cmd_score(args) -> int:
    path = MINED / f"{name(args)}.accuracy.xlsx"
    ws = load_workbook(path)["Выборка"]
    back = {v: k for k, v in VERDICT.items()}
    result = 0
    for section in ("Триаж", "Сопоставление"):
        good = bad = open_ = 0
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[1] != section:
                continue
            mark = (r[10] or "").strip().lower()
            mark = "" if mark == "не знаю" else mark  # falls back to Claude's verdict
            if mark == "да" or (not mark and back.get(r[7]) == "ok"):
                good += 1
            elif mark == "нет" or (not mark and back.get(r[7]) == "wrong"):
                bad += 1
            else:
                open_ += 1
        if not good + bad + open_:
            continue
        pct = 100 * good / (good + bad) if good + bad else 0
        ok = "цель достигнута" if pct >= TARGET and not open_ else "открыто" if open_ else "ниже цели"
        print(f"{section}: верно {good}, неверно {bad}, открыто {open_} — {pct:.1f} % ({ok})")
        result |= pct < TARGET
    return result


def main() -> int:
    setup_console()
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("cmd", choices=["build", "score"])
    ap.add_argument("--batch", default="pilot", help="партия или несколько через запятую")
    ap.add_argument("--name", help="имя выборки (файлы <name>.accuracy.*); по умолчанию — --batch")
    ap.add_argument("--match-only", action="store_true", help="только сопоставление, без триажа")
    ap.add_argument("--n", type=int, default=SAMPLE_MATCH, help="размер выборки сопоставления (не для пилота)")
    ap.add_argument("--marks", help="перенос отметок: файл со строками KEY=отметка|комментарий")
    args = ap.parse_args()
    return {"build": cmd_build, "score": cmd_score}[args.cmd](args)


if __name__ == "__main__":
    sys.exit(main())
